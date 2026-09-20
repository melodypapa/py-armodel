from pathlib import Path

import pytest
from openpyxl import load_workbook

from armodel.cli.os_config_export_cli import main


def test_cli_defaults_to_xlsx(monkeypatch, tmp_path: Path):
    output = tmp_path / "os.xlsx"
    monkeypatch.setattr("sys.argv", ["os-config-export", "tests/integration_tests/test_files/Os_ECUC.arxml", str(output)])

    main()

    assert output.exists()
    assert "OsTask" in load_workbook(output).sheetnames
    assert (tmp_path / "os_config_export.log").exists()


def test_cli_supports_yaml(monkeypatch, tmp_path: Path):
    output = tmp_path / "os.yaml"
    monkeypatch.setattr("sys.argv", ["os-config-export", "tests/integration_tests/test_files/Os_ECUC.arxml", str(output), "--format", "yaml"])

    main()

    assert output.exists()
    assert "name: OsApplication_QM" in output.read_text(encoding="utf-8")
    assert (tmp_path / "os_config_export.log").exists()


def test_cli_warning_flag_downgrades_unresolved_references(monkeypatch, tmp_path: Path, caplog):
    output = tmp_path / "os.xlsx"
    monkeypatch.setattr("sys.argv", ["os-config-export", "-w", "tests/integration_tests/test_files/Os_ECUC.arxml", str(output)])

    main()

    assert output.exists()


def test_cli_rejects_output_equal_to_input(monkeypatch, tmp_path: Path):
    input_file = tmp_path / "input.arxml"
    input_file.write_text("dummy")
    monkeypatch.setattr("sys.argv", ["os-config-export", str(input_file), str(input_file)])

    with pytest.raises(ValueError, match="must not overwrite the input file"):
        main()

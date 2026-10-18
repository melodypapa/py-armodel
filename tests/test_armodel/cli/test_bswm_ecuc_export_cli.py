from pathlib import Path

import pytest
from openpyxl import load_workbook

from armodel.cli.bswm_ecuc_export_cli import main

FIXTURE = Path(__file__).resolve().parents[1] / "parser" / "data" / "bswm_ecuc_minimal.arxml"


def test_cli_rejects_output_equal_to_input(monkeypatch, tmp_path: Path):
    input_file = tmp_path / "input.arxml"
    input_file.write_text("dummy")
    monkeypatch.setattr("sys.argv", ["bswm-ecuc-export", str(input_file), str(input_file)])

    with pytest.raises(ValueError, match="must not overwrite the input file"):
        main()


def test_cli_exports_the_fixture_in_warning_mode(monkeypatch, tmp_path: Path):
    output = tmp_path / "bswm.xlsx"
    monkeypatch.setattr("sys.argv", ["bswm-ecuc-export", "-w", str(FIXTURE), str(output)])

    main()

    assert load_workbook(output).sheetnames == ["BswMAction", "BswMActionList"]
    assert (tmp_path / "bswm_ecuc_export.log").exists()

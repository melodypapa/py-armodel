from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from armodel.data_models.ecuc import OsAlarm, OsApplication, OsIsr, OsOs, OsScheduleTable, OsScheduleTableExpiryPoint, OsTask
from armodel.report import OsConfigExporter, OsConfigModelMapper, OsConfigXlsxExporter, OsConfigYamlExporter


def test_write_yaml_uses_semantic_names(tmp_path: Path):
    os_os = OsOs().setName("Os")
    os_os.addOsApplication(OsApplication().setName("App"))
    os_os.addOsTask(OsTask().setName("Task").setOsTaskPriority(5))
    output = tmp_path / "os.yaml"

    OsConfigYamlExporter().export(os_os, output)

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    assert data["OsApplication"][0]["name"] == "App"
    assert data["OsTask"][0]["OsTaskPriority"] == 5


def test_write_yaml_omits_none_and_empty_list_values(tmp_path: Path):
    os_os = OsOs().setName("Os")
    os_os.addOsTask(OsTask().setName("Task").setOsTaskPriority(5))
    output = tmp_path / "os.yaml"

    OsConfigYamlExporter().export(os_os, output)

    task = yaml.safe_load(output.read_text(encoding="utf-8"))["OsTask"][0]
    assert task == {"name": "Task", "OsTaskPriority": 5}


def test_os_config_export_rejects_unsupported_extension(tmp_path: Path):
    with pytest.raises(ValueError, match="Unsupported OS configuration export format"):
        OsConfigExporter().export(OsOs(), tmp_path / "os.json")


def test_write_xlsx_uses_expected_sheets(tmp_path: Path):
    output = tmp_path / "os.xlsx"
    os_os = OsOs().setName("Os")
    os_os.addOsApplication(OsApplication().setName("App"))
    os_os.addOsTask(OsTask().setName("Task"))
    OsConfigXlsxExporter().export(os_os, output)

    workbook = load_workbook(output)
    assert "OsApplication" in workbook.sheetnames
    assert "OsTask" in workbook.sheetnames


def test_write_yaml_without_pyyaml_raises_actionable_error(tmp_path: Path, monkeypatch):
    import builtins

    real_import = builtins.__import__

    def blocked_import(name, *args, **kwargs):
        if name == "yaml":
            raise ImportError("No module named 'yaml'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked_import)
    os_os = OsOs().setName("Os")
    with pytest.raises(ImportError, match="pip install pyyaml"):
        OsConfigYamlExporter().export(os_os, tmp_path / "os.yaml")


def test_write_yaml_represents_relationships_by_name(tmp_path: Path):
    application = OsApplication().setName("App")
    task = OsTask().setName("Task")
    application.addOsAppTaskRef(task)
    task.addOsTaskAccessingApplication(application)
    application.setOsRestartTask(task)
    os_os = OsOs().setName("Os")
    os_os.addOsApplication(application)
    os_os.addOsTask(task)
    output = tmp_path / "os.yaml"

    OsConfigYamlExporter().export(os_os, output)

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    assert data["OsApplication"][0]["OsAppTaskRef"] == ["Task"]
    assert data["OsApplication"][0]["OsRestartTask"] == "Task"
    assert data["OsTask"][0]["OsTaskAccessingApplication"] == ["App"]


def test_write_yaml_keeps_flattened_and_reference_fields(tmp_path: Path):
    task = OsTask().setName("Task")
    task.addOsTaskAppModeRef("/Os/Os/OSDEFAULTAPPMODE")
    task.setOsTaskExecutionBudget(0.005)
    task.addOsTaskResourceLockBudget(0.001)
    task.addOsTaskResourceLockResourceRef("/Os/Os/Resource_Init")
    os_os = OsOs().setName("Os")
    os_os.addOsTask(task)

    output = tmp_path / "os.yaml"
    OsConfigYamlExporter().export(os_os, output)

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    exported = data["OsTask"][0]
    assert exported["OsTaskAppModeRef"] == ["/Os/Os/OSDEFAULTAPPMODE"]
    assert exported["OsTaskExecutionBudget"] == 0.005
    assert exported["OsTaskResourceLockBudget"] == [0.001]
    assert exported["OsTaskResourceLockResourceRef"] == ["/Os/Os/Resource_Init"]


def test_write_xlsx_rows_use_name_identity_and_exact_fields(tmp_path: Path):
    task = OsTask().setName("Task").setOsTaskPriority(5).setOsStacksize(1024)
    task.addOsTaskEventRef("/Os/Os/Event1")
    task.addOsTaskEventRef("/Os/Os/Event2")
    application = OsApplication().setName("App").setOsTrusted(False)
    application.addOsAppTaskRef(task)
    os_os = OsOs().setName("Os")
    os_os.addOsApplication(application)
    os_os.addOsTask(task)
    output = tmp_path / "os.xlsx"

    OsConfigXlsxExporter().export(os_os, output)

    workbook = load_workbook(output)
    task_sheet = workbook["OsTask"]
    headers = [cell.value for cell in task_sheet[1]]
    assert headers[0] == "name"
    assert "OsTaskPriority" in headers
    assert "OsStacksize" in headers
    rows = list(task_sheet.iter_rows(min_row=2, values_only=True))
    assert rows[0][headers.index("name")] == "Task"
    assert rows[0][headers.index("OsTaskPriority")] == 5
    event_cell = task_sheet.cell(row=2, column=headers.index("OsTaskEventRef") + 1)
    assert event_cell.value == "/Os/Os/Event1\n/Os/Os/Event2"
    assert event_cell.alignment.wrap_text is True

    application_sheet = workbook["OsApplication"]
    app_headers = [cell.value for cell in application_sheet[1]]
    app_rows = list(application_sheet.iter_rows(min_row=2, values_only=True))
    assert app_rows[0][app_headers.index("name")] == "App"
    assert app_rows[0][app_headers.index("OsTrusted")] is False
    assert app_rows[0][app_headers.index("OsAppTaskRef")] == "Task"


def test_write_yaml_exports_alarm_isr_and_schedule_table_sections(tmp_path: Path):
    alarm = OsAlarm().setName("Alarm1").setOsAlarmCounterRef("/Os/Os/HwCounter").setOsAlarmIncrementCounterRef("/Os/Os/Rte_Counter").setOsAlarmAlarmTime(1).setOsAlarmAutostartType("RELATIVE")
    isr = OsIsr().setName("CanIsr").setOsIsrCategory("CATEGORY_2").setOsIsrPeriod(0.005).setOsIsrExecutionBudget(0.001)
    isr.addOsIsrResourceLockBudget(0.0005)
    isr.addOsIsrResourceLockResourceRef("/Os/Os/OsStackResource")
    expiry_point1 = OsScheduleTableExpiryPoint().setOsScheduleTableExpiryPointOffset(2).setOsScheduleTableActivateTaskRef("/Os/Os/Rte_Time_Task")
    expiry_point2 = OsScheduleTableExpiryPoint().setOsScheduleTableExpiryPointOffset(5).setOsScheduleTableMaxShorten(1).setOsScheduleTableSetEventTaskRef("/Os/Os/Rte_Event_Task").setOsScheduleTableSetEventRef("/Os/Os/Rte_OSShutdownEvent")
    schedule_table = (
        OsScheduleTable()
        .setName("SystemScheduleTable")
        .setOsScheduleTableCounterRef("/Os/Os/HwCounter")
        .setOsScheduleTableDuration(10)
        .setOsScheduleTableRepeating(True)
        .setOsScheduleTableAutostartType("RELATIVE")
        .setOsScheduleTableSyncStrategy("IMPLICIT")
    )
    schedule_table.addOsScheduleTableExpiryPoint(expiry_point1)
    schedule_table.addOsScheduleTableExpiryPoint(expiry_point2)
    os_os = OsOs().setName("Os")
    os_os.addOsAlarm(alarm)
    os_os.addOsIsr(isr)
    os_os.addOsScheduleTable(schedule_table)
    output = tmp_path / "os.yaml"

    OsConfigYamlExporter().export(os_os, output)

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    assert set(data) == {"OsApplication", "OsTask", "OsAlarm", "OsIsr", "OsScheduleTable"}
    assert data["OsAlarm"][0] == {
        "name": "Alarm1",
        "OsAlarmCounterRef": "/Os/Os/HwCounter",
        "OsAlarmIncrementCounterRef": "/Os/Os/Rte_Counter",
        "OsAlarmAlarmTime": 1,
        "OsAlarmAutostartType": "RELATIVE",
    }
    assert data["OsIsr"][0] == {
        "name": "CanIsr",
        "OsIsrCategory": "CATEGORY_2",
        "OsIsrPeriod": 0.005,
        "OsIsrExecutionBudget": 0.001,
        "OsIsrResourceLockBudget": [0.0005],
        "OsIsrResourceLockResourceRef": ["/Os/Os/OsStackResource"],
    }
    assert data["OsScheduleTable"][0] == {
        "name": "SystemScheduleTable",
        "OsScheduleTableCounterRef": "/Os/Os/HwCounter",
        "OsScheduleTableDuration": 10,
        "OsScheduleTableRepeating": True,
        "OsScheduleTableAutostartType": "RELATIVE",
        "OsScheduleTableSyncStrategy": "IMPLICIT",
        "OsScheduleTableExpiryPoint": [
            {"OsScheduleTblExpPointOffset": 2, "OsScheduleTableActivateTaskRef": "/Os/Os/Rte_Time_Task"},
            {
                "OsScheduleTblExpPointOffset": 5,
                "OsScheduleTableMaxShorten": 1,
                "OsScheduleTableSetEventTaskRef": "/Os/Os/Rte_Event_Task",
                "OsScheduleTableSetEventRef": "/Os/Os/Rte_OSShutdownEvent",
            },
        ],
    }


def test_mapper_to_dict_contains_five_sections_with_defaults():
    data = OsConfigModelMapper().to_dict(OsOs().setName("Os"))

    assert set(data) == {"OsApplication", "OsTask", "OsAlarm", "OsIsr", "OsScheduleTable"}
    assert data["OsApplication"] == []
    assert data["OsTask"] == []
    assert data["OsAlarm"] == []
    assert data["OsIsr"] == []
    assert data["OsScheduleTable"] == []


def test_write_xlsx_uses_correct_headers_for_empty_new_sections(tmp_path: Path):
    output = tmp_path / "os.xlsx"

    OsConfigXlsxExporter().export(OsOs().setName("Os"), output)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["OsApplication", "OsTask", "OsAlarm", "OsIsr", "OsScheduleTable"]
    isr_headers = [cell.value for cell in workbook["OsIsr"][1]]
    assert isr_headers[0] == "name"
    assert "OsIsrCategory" in isr_headers
    table_headers = [cell.value for cell in workbook["OsScheduleTable"][1]]
    assert table_headers[0] == "name"
    assert "OsScheduleTableCounterRef" in table_headers


def test_write_xlsx_formats_expiry_point_dicts_readably(tmp_path: Path):
    expiry_point = OsScheduleTableExpiryPoint().setOsScheduleTableExpiryPointOffset(2).setOsScheduleTableActivateTaskRef("/Os/Os/Task1")
    schedule_table = OsScheduleTable().setName("Table1").setOsScheduleTableDuration(10)
    schedule_table.addOsScheduleTableExpiryPoint(expiry_point)
    os_os = OsOs().setName("Os")
    os_os.addOsScheduleTable(schedule_table)
    output = tmp_path / "os.xlsx"

    OsConfigXlsxExporter().export(os_os, output)

    workbook = load_workbook(output)
    sheet = workbook["OsScheduleTable"]
    headers = [cell.value for cell in sheet[1]]
    cell = sheet.cell(row=2, column=headers.index("OsScheduleTableExpiryPoint") + 1)
    assert cell.value == "OsScheduleTblExpPointOffset=2,OsScheduleTableActivateTaskRef=/Os/Os/Task1"

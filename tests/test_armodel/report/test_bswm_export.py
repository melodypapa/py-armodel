from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from armodel.data_models.ecuc import BswmAction, BswmActionList, BswmActionListItem, BswmModule
from armodel.report import BswmConfigExporter, BswmConfigModelMapper, BswmConfigXlsxExporter


def build_module() -> BswmModule:
    module = BswmModule().setName("BswM")
    module.addAction(BswmAction().setName("AC_UserCallout").setActionType("BswMUserCallout").addParameter("BswMUserCalloutFunction", "App_OnStartup()"))
    module.addAction(
        BswmAction()
        .setName("AC_PduGroupSwitch")
        .setActionType("BswMPduGroupSwitch")
        .addParameter("BswMPduGroupSwitchReinit", True)
        .addParameter("BswMEnabledPduGroupRef", "/Com/Tx")
        .addParameter("BswMEnabledPduGroupRef", "/Com/Rx")
    )
    action_list = BswmActionList().setName("AL_Startup").setExecution("BSWM_TRIGGER").setPriority(5)
    action_list.addItem(BswmActionListItem().setName("AI_Second").setIndex(20).setAbortOnFail(False).setActionRef("/BswM/C/AC_PduGroupSwitch").setActionRefType("BswMAction"))
    action_list.addItem(
        BswmActionListItem().setName("AI_First").setIndex(10).setAbortOnFail(True).setReportFailRuntimeErrorId(200).setActionRef("/BswM/C/AC_UserCallout").setActionRefType("BswMAction")
    )
    module.addActionList(action_list)
    module.addActionList(BswmActionList().setName("AL_Empty").setExecution("BSWM_CONDITION"))
    return module


def test_mapper_flattens_actions_with_one_column_per_parameter():
    rows = BswmConfigModelMapper().to_dict(build_module())["BswMAction"]

    assert rows[0] == {"name": "AC_UserCallout", "BswMActionType": "BswMUserCallout", "BswMUserCalloutFunction": "App_OnStartup()"}
    assert rows[1]["BswMEnabledPduGroupRef"] == ["/Com/Tx", "/Com/Rx"]
    assert rows[1]["BswMPduGroupSwitchReinit"] is True


def test_mapper_writes_one_action_list_row_per_item_sorted_by_index():
    rows = BswmConfigModelMapper().to_dict(build_module())["BswMActionList"]

    assert [(row["name"], row["BswMActionListItem"], row["BswMActionListItemIndex"], row["BswMActionListItemRef"]) for row in rows] == [
        ("AL_Startup", "AI_First", 10, "AC_UserCallout"),
        ("AL_Startup", "AI_Second", 20, "AC_PduGroupSwitch"),
        ("AL_Empty", None, None, None),
    ]


def test_xlsx_has_action_sheets_first_with_union_headers(tmp_path: Path):
    output = tmp_path / "bswm.xlsx"

    BswmConfigXlsxExporter().export(build_module(), output)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["BswMAction", "BswMActionList"]
    action_sheet = workbook["BswMAction"]
    assert [cell.value for cell in action_sheet[1]] == ["name", "BswMActionType", "BswMUserCalloutFunction", "BswMPduGroupSwitchReinit", "BswMEnabledPduGroupRef"]
    assert action_sheet.cell(row=3, column=5).value == "/Com/Tx\n/Com/Rx"
    list_sheet = workbook["BswMActionList"]
    assert [cell.value for cell in list_sheet[1]] == [
        "name",
        "BswMActionListExecution",
        "BswMActionListPriority",
        "BswMActionListItem",
        "BswMActionListItemIndex",
        "BswMAbortOnFail",
        "BswMReportFailRuntimeErrorId",
        "BswMActionListItemRef",
        "BswMActionListItemRefType",
    ]
    assert (list_sheet.cell(row=2, column=3).value, list_sheet.cell(row=2, column=4).value, list_sheet.cell(row=2, column=7).value) == (5, "AI_First", 200)


def test_xlsx_writes_headers_for_an_empty_module(tmp_path: Path):
    output = tmp_path / "bswm.xlsx"

    BswmConfigXlsxExporter().export(BswmModule().setName("BswM"), output)

    workbook = load_workbook(output)
    assert [cell.value for cell in workbook["BswMAction"][1]] == ["name", "BswMActionType"]


def test_yaml_export_omits_empty_values(tmp_path: Path):
    output = tmp_path / "bswm.yaml"

    BswmConfigExporter().export(build_module(), output)

    data = yaml.safe_load(output.read_text(encoding="utf-8"))
    assert data["BswMAction"][0] == {"name": "AC_UserCallout", "BswMActionType": "BswMUserCallout", "BswMUserCalloutFunction": "App_OnStartup()"}
    assert data["BswMActionList"][2] == {"name": "AL_Empty", "BswMActionListExecution": "BSWM_CONDITION"}


def test_exporter_rejects_unsupported_extension(tmp_path: Path):
    with pytest.raises(ValueError, match="Unsupported BswM configuration export format: json"):
        BswmConfigExporter().export(build_module(), tmp_path / "bswm.json")

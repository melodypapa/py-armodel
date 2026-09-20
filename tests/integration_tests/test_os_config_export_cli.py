from pathlib import Path

import yaml
from openpyxl import load_workbook

from armodel.cli.os_config_export_cli import main

OS_ECUC_FILE = Path(__file__).parent / "test_files" / "Os_ECUC.arxml"
EXPECTED_YAML_FILE = Path(__file__).parent / "test_files" / "Os_ECUC.yaml"

APPLICATION_FIELDS = {
    "name",
    "OsTrusted",
    "OsTrustedApplicationDelayTimingViolationCall",
    "OsTrustedApplicationWithProtection",
    "OsAppAlarmRef",
    "OsAppCounterRef",
    "OsAppEcucPartitionRef",
    "OsAppIsrRef",
    "OsAppScheduleTableRef",
    "OsAppTaskRef",
    "OsMemoryMappingCodeLocationRef",
    "OsRestartTask",
    "OsAppStartupHook",
    "OsAppErrorHook",
    "OsAppShutdownHook",
    "OsTrustedFunctionName",
    "ApplicationState",
}

TASK_FIELDS = {
    "name",
    "OsTaskActivation",
    "OsTaskPeriod",
    "OsTaskPriority",
    "OsTaskSchedule",
    "OsStacksize",
    "OsMemoryMappingCodeLocationRef",
    "OsTaskAccessingApplication",
    "OsTaskEventRef",
    "OsTaskResourceRef",
    "OsTaskAppModeRef",
    "OsTaskAllInterruptLockBudget",
    "OsTaskExecutionBudget",
    "OsTaskOsInterruptLockBudget",
    "OsTaskTimeFrame",
    "OsTaskResourceLockBudget",
    "OsTaskResourceLockResourceRef",
}


def test_os_config_export_cli_exports_real_os_ecuc_to_xlsx(monkeypatch, tmp_path: Path):
    output = tmp_path / "os.xlsx"
    monkeypatch.setattr("sys.argv", ["os-config-export", str(OS_ECUC_FILE), str(output)])

    main()

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["OsApplication", "OsTask"]
    assert workbook["OsApplication"]["A2"].value == "OsApplication_QM"
    assert workbook["OsTask"]["A2"].value == "Init_Task"
    assert (tmp_path / "os_config_export.log").exists()


def test_os_config_export_cli_exports_real_os_ecuc_to_yaml(monkeypatch, tmp_path: Path):
    output = tmp_path / "os.yaml"
    monkeypatch.setattr("sys.argv", ["os-config-export", str(OS_ECUC_FILE), str(output), "--format", "yaml"])

    main()

    content = output.read_text(encoding="utf-8")
    assert "name: OsApplication_QM" in content
    assert "name: Init_Task" in content
    assert (tmp_path / "os_config_export.log").exists()


def test_os_config_export_cli_matches_complete_yaml_fixture(monkeypatch, tmp_path: Path):
    output = tmp_path / "os.yaml"
    monkeypatch.setattr("sys.argv", ["os-config-export", str(OS_ECUC_FILE), str(output), "--format", "yaml"])

    main()

    generated = yaml.safe_load(output.read_text(encoding="utf-8"))
    expected = yaml.safe_load(EXPECTED_YAML_FILE.read_text(encoding="utf-8"))

    assert generated == expected
    assert set(generated) == {"OsApplication", "OsTask"}
    assert generated["OsApplication"]
    assert generated["OsTask"]
    assert all(set(application) == APPLICATION_FIELDS for application in generated["OsApplication"])
    assert all(set(task) == TASK_FIELDS for task in generated["OsTask"])
